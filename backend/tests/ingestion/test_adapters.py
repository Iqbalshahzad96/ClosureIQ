"""
Tests for Source Adapters.

Covers:
- Enquest Ledger adapter (OOXML detection, dynamic headers, missing Remarks,
  row-role classification, date/Decimal parsing, dual-sided flagging)
- Enquest Trial Balance adapter
- Generic Bank Statement (CSV and Excel)
- Generic AP Invoice (CSV and Excel)
- Generic Fixed Asset Register (CSV and Excel)
- Adapter Registry resolution
"""

import io
from datetime import datetime
from decimal import Decimal
from typing import List

import openpyxl
import pytest

from app.ingestion.adapters.enquest_ledger import EnquestLedgerAdapter
from app.ingestion.adapters.enquest_tb import EnquestTBAdapter
from app.ingestion.adapters.generic_ap_invoice import GenericAPInvoiceAdapter
from app.ingestion.adapters.generic_bank import GenericBankStatementAdapter
from app.ingestion.adapters.generic_fixed_asset import GenericFixedAssetAdapter
from app.ingestion.adapters.registry import AdapterRegistry
from app.ingestion.base import CanonicalRecordPayload, RawSourceRow, RowRole


# ===========================================================================
# Helpers — generate in-memory OOXML workbooks
# ===========================================================================


def _make_xlsx_bytes(rows, sheet_name="Sheet1"):
    """Create a minimal .xlsx (OOXML) workbook from a list of row-tuples."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    for row in rows:
        ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def _make_csv_bytes(rows):
    """Create CSV bytes from a list of row-lists."""
    import csv
    buf = io.StringIO()
    writer = csv.writer(buf)
    for row in rows:
        writer.writerow(row)
    return buf.getvalue().encode("utf-8")


# ===========================================================================
# Enquest Ledger Adapter
# ===========================================================================


class TestEnquestLedgerAdapter:
    """Tests for the Enquest ERP Ledger adapter."""

    def _build_ledger_workbook(self, include_remarks=True, dual_sided=False):
        """Build a representative Enquest ledger workbook."""
        rows = [
            ("Ledger Inquiry",),
            ("Ledger: DIAMOND TRUST BANK KSHS",),
            ("Period: 01-Jan-2025 to 31-Dec-2025",),
        ]
        if include_remarks:
            rows.append((
                "Voucher Date", "Particulars", "Voucher Type", "Voucher No",
                "Ref No", "Cheque No", "Curr.", "Ex.Rate",
                "Debit Amount", "Credit Amount", "Balance Amount",
                "Is Bank Reconciliated", "CUIN/FDA", "Remarks",
            ))
        else:
            rows.append((
                "Voucher Date", "Particulars", "Voucher Type", "Voucher No",
                "Ref No", "Cheque No", "Curr.", "Ex.Rate",
                "Debit Amount", "Credit Amount", "Balance Amount",
                "Is Bank Reconciliated", "CUIN/FDA",
            ))

        # Opening balance row
        rows.append((
            None, "Opening Balance", None, None,
            None, None, "KES", 1.0,
            0, 0, 100000.50,
            None, None,
        ) + (("",) if include_remarks else ()))

        # Transaction row
        debit = 50000.75
        credit = 0 if not dual_sided else 10000.25
        rows.append((
            datetime(2025, 1, 4), "Customer Cash Receipt", "Receipt", "V001",
            "REF-100", "CHQ-001", "KES", 1.0,
            debit, credit, 150001.25,
            "Yes", "CUIN-001",
        ) + (("Payment for services",) if include_remarks else ()))

        # Total row
        rows.append((
            None, "Total", None, None,
            None, None, None, None,
            50000.75, 0, 150001.25,
            None, None,
        ) + (("",) if include_remarks else ()))

        return _make_xlsx_bytes(rows)

    def test_can_handle_valid_ooxml(self):
        adapter = EnquestLedgerAdapter()
        content = self._build_ledger_workbook()
        assert adapter.can_handle(content, "DIAMOND TRUST BANK KSHS.xls")

    def test_can_handle_rejects_csv(self):
        adapter = EnquestLedgerAdapter()
        csv_bytes = b"date,amount\n2025-01-01,100"
        assert not adapter.can_handle(csv_bytes, "test.csv")

    def test_dynamic_header_detection(self):
        adapter = EnquestLedgerAdapter()
        content = self._build_ledger_workbook()
        rows = list(adapter.parse_file(content, "DIAMOND TRUST BANK KSHS.xls"))
        header_rows = [r for r in rows if r.role == RowRole.HEADER]
        assert len(header_rows) == 1

    def test_row_role_classification(self):
        adapter = EnquestLedgerAdapter()
        content = self._build_ledger_workbook()
        rows = list(adapter.parse_file(content, "test.xls"))
        roles = [r.role for r in rows]
        assert RowRole.TITLE in roles
        assert RowRole.HEADER in roles
        assert RowRole.OPENING_BALANCE in roles
        assert RowRole.TRANSACTION in roles
        assert RowRole.TOTAL in roles

    def test_missing_remarks_column(self):
        adapter = EnquestLedgerAdapter()
        content = self._build_ledger_workbook(include_remarks=False)
        rows = list(adapter.parse_file(content, "test.xls"))
        tx_rows = [r for r in rows if r.role == RowRole.TRANSACTION]
        assert len(tx_rows) >= 1
        # Should still parse without error
        payloads = adapter.map_to_canonical(rows)
        assert len(payloads) >= 1

    def test_decimal_parsing(self):
        adapter = EnquestLedgerAdapter()
        content = self._build_ledger_workbook()
        rows = list(adapter.parse_file(content, "test.xls"))
        payloads = adapter.map_to_canonical(rows)
        assert len(payloads) >= 1
        line = payloads[0].data["lines"][0]
        assert isinstance(line["debit_amount"], Decimal)
        assert line["debit_amount"] == Decimal("50000.7500")

    def test_dual_sided_posting_flag(self):
        adapter = EnquestLedgerAdapter()
        content = self._build_ledger_workbook(dual_sided=True)
        rows = list(adapter.parse_file(content, "test.xls"))
        payloads = adapter.map_to_canonical(rows)
        assert len(payloads) >= 1
        assert "FLAG_DUAL_SIDED_POSTING" in payloads[0].flags

    def test_ooxml_incorrectly_named_xls(self):
        """OOXML content with .xls extension should still be parsed."""
        adapter = EnquestLedgerAdapter()
        content = self._build_ledger_workbook()
        # Filename ends with .xls but content is OOXML (PK header)
        assert content[:4] == b"PK\x03\x04"
        assert adapter.can_handle(content, "ledger.xls")
        rows = list(adapter.parse_file(content, "ledger.xls"))
        assert len(rows) > 0

    def test_excludes_non_transaction_rows_from_canonical(self):
        adapter = EnquestLedgerAdapter()
        content = self._build_ledger_workbook()
        rows = list(adapter.parse_file(content, "test.xls"))
        payloads = adapter.map_to_canonical(rows)
        # Only TRANSACTION rows should produce payloads
        for p in payloads:
            assert p.entity_type == "JOURNAL_ENTRY"

    def test_preserves_is_bank_reconciliated_as_metadata(self):
        adapter = EnquestLedgerAdapter()
        content = self._build_ledger_workbook()
        rows = list(adapter.parse_file(content, "test.xls"))
        payloads = adapter.map_to_canonical(rows)
        assert len(payloads) >= 1
        dims = payloads[0].data["lines"][0].get("dimensions_json", {})
        assert "source_is_bank_reconciled" in dims

    def test_lineage_preserved(self):
        adapter = EnquestLedgerAdapter()
        content = self._build_ledger_workbook()
        rows = list(adapter.parse_file(content, "test.xls"))
        payloads = adapter.map_to_canonical(rows)
        assert len(payloads) >= 1
        assert payloads[0].raw_lineage.get("filename") is not None
        assert payloads[0].raw_lineage.get("row_number") is not None


# ===========================================================================
# Enquest Trial Balance Adapter
# ===========================================================================


class TestEnquestTBAdapter:

    def _build_tb_workbook(self):
        rows = [
            ("Enquest Trial Balance — 2025",),
            ("Account", "Opening Balance", "Period Debit",
             "Period Credit", "Closing Balance"),
            ("Cash at Bank", 100000, 50000, 30000, 120000),
            ("Accounts Payable", 200000, 10000, 80000, 270000),
            ("Grand Total", 300000, 60000, 110000, 390000),
        ]
        return _make_xlsx_bytes(rows)

    def test_can_handle_tb(self):
        adapter = EnquestTBAdapter()
        content = self._build_tb_workbook()
        assert adapter.can_handle(content, "Enquest TB 2025.xls")

    def test_parse_and_map(self):
        adapter = EnquestTBAdapter()
        content = self._build_tb_workbook()
        rows = list(adapter.parse_file(content, "Enquest TB 2025.xls"))
        payloads = adapter.map_to_canonical(rows)
        # Should have 2 account rows (not grand total)
        assert len(payloads) == 2
        assert payloads[0].entity_type == "TRIAL_BALANCE"
        assert payloads[0].data["account_name_raw"] == "Cash at Bank"
        assert payloads[0].data["opening_balance"] == Decimal("100000.0000")

    def test_tb_quarantines_malformed(self):
        adapter = EnquestTBAdapter()
        bad_bytes = b"This is not a workbook"
        # A malformed file is an explicit failure, not a successful empty import.
        with pytest.raises(ValueError):
            list(adapter.parse_file(bad_bytes, "bad.xls"))


# ===========================================================================
# Generic Bank Statement Adapter
# ===========================================================================


class TestGenericBankStatementAdapter:

    def test_csv_with_signed_amount(self):
        adapter = GenericBankStatementAdapter()
        csv = _make_csv_bytes([
            ["Booking Date", "Description", "Amount", "Balance", "Reference"],
            ["2025-01-05", "Wire Deposit", "50000.00", "150000.00", "FT25005"],
            ["2025-01-06", "Payment Out", "-20000.00", "130000.00", "FT25006"],
        ])
        assert adapter.can_handle(csv, "bank_statement.csv")
        rows = list(adapter.parse_file(csv, "bank_statement.csv"))
        payloads = adapter.map_to_canonical(rows)
        assert len(payloads) == 2
        assert payloads[0].entity_type == "BANK_TRANSACTION"
        assert payloads[0].data["amount"] == Decimal("50000.0000")
        assert payloads[1].data["amount"] == Decimal("-20000.0000")

    def test_csv_with_deposit_withdrawal_columns(self):
        adapter = GenericBankStatementAdapter()
        csv = _make_csv_bytes([
            ["Date", "Description", "Deposit", "Withdrawal", "Balance"],
            ["2025-01-05", "Incoming Wire", "50000", "", "150000"],
            ["2025-01-06", "Cheque Cleared", "", "20000", "130000"],
        ])
        rows = list(adapter.parse_file(csv, "statement.csv"))
        payloads = adapter.map_to_canonical(rows)
        assert len(payloads) == 2
        assert payloads[0].data["amount"] == Decimal("50000.0000")
        assert payloads[1].data["amount"] == Decimal("-20000.0000")

    def test_excel_bank_statement(self):
        adapter = GenericBankStatementAdapter()
        content = _make_xlsx_bytes([
            ("Booking Date", "Description", "Amount", "Balance", "Bank Reference"),
            (datetime(2025, 1, 5), "Wire", 50000, 150000, "FT001"),
        ])
        assert adapter.can_handle(content, "bank_statement.xlsx")
        rows = list(adapter.parse_file(content, "bank_statement.xlsx"))
        payloads = adapter.map_to_canonical(rows)
        assert len(payloads) == 1


# ===========================================================================
# Generic AP Invoice Adapter
# ===========================================================================


class TestGenericAPInvoiceAdapter:

    def test_csv_ap_invoice(self):
        adapter = GenericAPInvoiceAdapter()
        csv = _make_csv_bytes([
            ["Vendor Name", "Invoice Number", "Invoice Date", "Due Date",
             "Total Amount", "Tax Amount", "Status"],
            ["Acme Corp", "INV-001", "2025-01-15", "2025-02-15",
             "116000.00", "16000.00", "OPEN"],
        ])
        assert adapter.can_handle(csv, "ap_invoices.csv")
        rows = list(adapter.parse_file(csv, "ap_invoices.csv"))
        payloads = adapter.map_to_canonical(rows)
        assert len(payloads) == 1
        assert payloads[0].entity_type == "AP_INVOICE"
        assert payloads[0].data["vendor_name"] == "Acme Corp"
        assert payloads[0].data["total_amount"] == Decimal("116000.0000")

    def test_excel_ap_invoice(self):
        adapter = GenericAPInvoiceAdapter()
        content = _make_xlsx_bytes([
            ("Vendor Name", "Invoice Number", "Invoice Date",
             "Total Amount", "Tax Amount"),
            ("Office Landlord", "INV-RENT-01", datetime(2025, 1, 1),
             116000, 16000),
        ])
        assert adapter.can_handle(content, "ap_invoice_register.xlsx")
        rows = list(adapter.parse_file(content, "ap_invoice_register.xlsx"))
        payloads = adapter.map_to_canonical(rows)
        assert len(payloads) == 1


# ===========================================================================
# Generic Fixed Asset Adapter
# ===========================================================================


class TestGenericFixedAssetAdapter:

    def test_csv_fixed_asset(self):
        adapter = GenericFixedAssetAdapter()
        csv = _make_csv_bytes([
            ["Asset Code", "Asset Name", "Category", "Acquisition Date",
             "Acquisition Cost", "Salvage Value", "Useful Life",
             "Accumulated Depreciation", "Book Value"],
            ["FA-001", "MacBook Pro", "COMPUTERS", "2025-01-10",
             "350000", "35000", "3", "8750", "341250"],
        ])
        assert adapter.can_handle(csv, "fixed_asset_register.csv")
        rows = list(adapter.parse_file(csv, "fixed_asset_register.csv"))
        payloads = adapter.map_to_canonical(rows, {"useful_life_unit": "years"})
        assert len(payloads) == 1
        assert payloads[0].entity_type == "FIXED_ASSET"
        assert payloads[0].data["asset_name"] == "MacBook Pro"
        # Explicit source units, never a heuristic based on the numeric magnitude.
        assert payloads[0].data["useful_life_months"] == 36

    def test_excel_fixed_asset(self):
        adapter = GenericFixedAssetAdapter()
        content = _make_xlsx_bytes([
            ("Asset Code", "Asset Name", "Category", "Acquisition Date",
             "Acquisition Cost", "Accumulated Depreciation"),
            ("FA-TRUCK-01", "Scania Truck", "MOTOR_VEHICLES",
             datetime(2023, 6, 1), 12000000, 2250000),
        ])
        assert adapter.can_handle(content, "asset_register.xlsx")
        rows = list(adapter.parse_file(content, "asset_register.xlsx"))
        payloads = adapter.map_to_canonical(rows)
        assert len(payloads) == 1
        assert payloads[0].data["acquisition_cost"] == Decimal("12000000.0000")


# ===========================================================================
# Adapter Registry
# ===========================================================================


class TestAdapterRegistry:

    def test_registry_resolves_enquest_ledger(self):
        registry = AdapterRegistry()
        adapter = EnquestLedgerAdapter()
        content = TestEnquestLedgerAdapter()._build_ledger_workbook()
        resolved = registry.resolve_adapter(content, "DIAMOND TRUST BANK KSHS.xls")
        assert resolved is not None
        assert resolved.get_adapter_key() == "enquest_ledger"

    def test_registry_resolves_bank_csv(self):
        registry = AdapterRegistry()
        csv = _make_csv_bytes([
            ["Date", "Description", "Amount", "Balance"],
            ["2025-01-01", "Test", "100", "100"],
        ])
        resolved = registry.resolve_adapter(csv, "bank_statement.csv")
        assert resolved is not None
        assert resolved.get_adapter_key() == "generic_bank"

    def test_registry_explicit_key(self):
        registry = AdapterRegistry()
        resolved = registry.resolve_adapter(b"", "anything", explicit_adapter_key="generic_ap_invoice")
        assert resolved is not None
        assert resolved.get_adapter_key() == "generic_ap_invoice"

    def test_registry_returns_none_for_unknown(self):
        registry = AdapterRegistry()
        resolved = registry.resolve_adapter(b"random unknown bytes", "unknown_file.dat")
        assert resolved is None

    def test_registry_get_by_key(self):
        registry = AdapterRegistry()
        assert registry.get_by_key("enquest_ledger") is not None
        assert registry.get_by_key("nonexistent") is None


def test_reordered_headers_and_later_sheet():
    wb=openpyxl.Workbook(); wb.active.title='Cover'; wb.active.append(['Synthetic cover'])
    ws=wb.create_sheet('Transactions')
    ws.append(['Credit Amount',' Voucher   Date ','Voucher No','Debit Amount','Voucher Type','Particulars','Curr.','Ex.Rate'])
    ws.append([0,45659,'V1','12.3456','Receipt','Demo','KES',1])
    buf=io.BytesIO(); wb.save(buf)
    adapter=EnquestLedgerAdapter()
    assert adapter.can_handle(buf.getvalue(),'random.xls')
    p=adapter.map_to_canonical(list(adapter.parse_file(buf.getvalue(),'random.xls')))[0]
    assert p.sheet_name=='Transactions' and p.source_row_number==2
    assert p.data['entry_date']==datetime(2025,1,2)
    assert p.data['lines'][0]['debit_amount']==Decimal('12.3456')


def test_opening_voucher_subtotal_footer_classification():
    adapter=EnquestLedgerAdapter()
    assert adapter._classify_row({'voucher_no':'Opening','debit_amount':10})==RowRole.OPENING_BALANCE
    assert adapter._classify_row({'particulars':'Subtotal'})==RowRole.SUBTOTAL
    assert adapter._classify_row({'particulars':'Printed by system'})==RowRole.FOOTER
    assert adapter._classify_row({'debit_amount':10,'credit_amount':5})==RowRole.TOTAL
    assert adapter._classify_row({'particulars':'Total Logistics vendor','voucher_no':'V1'})==RowRole.TRANSACTION


def test_invalid_source_values_survive_for_pipeline():
    adapter=EnquestLedgerAdapter()
    raw=RawSourceRow(2,'S',{'voucher_date':'bad','voucher_no':'V1','debit_amount':'bad','credit_amount':None},
                     metadata={'filename':'demo.xls','account_name_raw':'Demo'})
    p=adapter.map_to_canonical([raw])[0]
    assert p.data['entry_date']=='bad'
    assert p.data['lines'][0]['debit_amount']=='bad'
    assert p.data['lines'][0]['credit_amount'] is None
    assert p.raw_lineage['raw_values']['voucher_date']=='bad'


def test_bank_aliases_do_not_confuse_value_date_or_withdrawal_amount():
    adapter=GenericBankStatementAdapter()
    data=_make_csv_bytes([['Value Date','Booking Date','Withdrawal','Deposit','Balance'],
                          ['2025-01-03','2025-01-01','25','','100']])
    p=adapter.map_to_canonical(list(adapter.parse_file(data,'x.csv')))[0]
    assert p.data['booking_date']==datetime(2025,1,1)
    assert p.data['value_date']==datetime(2025,1,3)
    assert p.data['amount']==-25


def test_bank_missing_amount_and_dual_amount_not_silently_zeroed():
    adapter=GenericBankStatementAdapter()
    assert adapter._calculate_signed_amount({'deposit':None,'withdrawal':None}) is None
    assert adapter._calculate_signed_amount({'deposit':'10','withdrawal':'5'})=='AMBIGUOUS_DUAL_BANK_AMOUNT'


def test_ap_exact_aliases_and_explicit_zero_values():
    adapter=GenericAPInvoiceAdapter()
    content=_make_csv_bytes([['Vendor Code','Vendor Name','Invoice Number','Tax Amount','Subtotal Amount','Total Amount','Paid Amount','Outstanding Amount'],
                            ['VC','Vendor','INV','20','100','120','120','0']])
    p=adapter.map_to_canonical(list(adapter.parse_file(content,'ap.csv')))[0]
    assert p.data['vendor_name']=='Vendor'
    assert p.data['total_amount']==120 and p.data['outstanding_amount']==0


def test_life_months_never_heuristically_multiplied():
    adapter=GenericFixedAssetAdapter()
    data=_make_csv_bytes([['Asset Name','Acquisition Cost','Useful Life Months'],['Demo','10','6']])
    p=adapter.map_to_canonical(list(adapter.parse_file(data,'a.csv')))[0]
    assert p.data['useful_life_months']=='6'
    assert p.data['salvage_value'] is None and p.data['acquisition_date'] is None


def test_trial_balance_period_not_invented():
    adapter=EnquestTBAdapter(); data=TestEnquestTBAdapter()._build_tb_workbook()
    p=adapter.map_to_canonical(list(adapter.parse_file(data,'tb.xls')))[0]
    assert p.data['fiscal_period'] is None
    assert p.raw_lineage['sheet_name']=='Sheet1'


def test_unknown_explicit_adapter_does_not_fall_back():
    registry=AdapterRegistry()
    data=_make_csv_bytes([['Date','Amount'],['2025-01-01','1']])
    assert registry.resolve_adapter(data,'bank.csv','typo') is None


def test_tb_hierarchy_controls_are_not_account_records():
    data=_make_xlsx_bytes([['Account','Opening','Debit','Credit','Closing'],
                          ['All Account',0,10,10,0],['Group',5,10,2,13],
                          ['    Leaf',5,10,2,13]])
    adapter=EnquestTBAdapter()
    rows=list(adapter.parse_file(data,'tb.xls'))
    assert [r.role for r in rows]==[RowRole.HEADER,RowRole.TOTAL,RowRole.SUBTOTAL,RowRole.TRANSACTION]
    assert len(adapter.map_to_canonical(rows))==1


def test_excel_1904_epoch_for_serial_ledger_dates():
    from openpyxl.utils.datetime import MAC_EPOCH
    workbook=openpyxl.Workbook(); workbook.epoch=MAC_EPOCH
    workbook.active.append(['Voucher Date','Voucher Type','Voucher No','Debit Amount','Credit Amount'])
    workbook.active.append([1,'Journal','V1',1,0])
    output=io.BytesIO(); workbook.save(output)
    adapter=EnquestLedgerAdapter()
    p=adapter.map_to_canonical(list(adapter.parse_file(output.getvalue(),'ledger.xls')))[0]
    assert p.data['entry_date']==datetime(1904,1,2)
